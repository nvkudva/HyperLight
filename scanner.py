from concurrent.futures import FIRST_COMPLETED, Future, ThreadPoolExecutor, wait
import datetime
import email
import email.policy
import functools
import hashlib
import html
import os
import re
from pathlib import Path
from typing import Any, Callable, Dict, Generator, List, Optional, Tuple
import docx
import pypdf
import pypdfium2 as pdfium
import embedder
import storage

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tiff"}
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json"}
CODE_EXTENSIONS = {".py", ".js", ".ts", ".swift", ".rs", ".go", ".sh", ".yaml", ".yml", ".html", ".css"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx"} | TEXT_EXTENSIONS | CODE_EXTENSIONS
CONTACT_EXTENSIONS = {".vcf"}
CALENDAR_EXTENSIONS = {".ics"}
BOOKMARK_EXTENSIONS = {".webloc", ".url"}
EMAIL_EXTENSIONS = {".emlx", ".eml"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".m4v", ".webm", ".mkv", ".avi"}
AUDIO_EXTENSIONS = {".mp3", ".m4a", ".wav", ".aac", ".flac", ".aiff", ".ogg"}
ALL_EXTENSIONS = (
    IMAGE_EXTENSIONS | DOCUMENT_EXTENSIONS | CONTACT_EXTENSIONS | CALENDAR_EXTENSIONS
    | BOOKMARK_EXTENSIONS | EMAIL_EXTENSIONS | VIDEO_EXTENSIONS | AUDIO_EXTENSIONS
)

IGNORED_DIRS = {
    ".git", ".venv", "venv", "node_modules", "__pycache__", ".cache",
    "build", "dist", ".Trash", ".idea", ".vscode", "Library", ".npm",
}

def collect_supported_files(target_path: Path) -> List[Tuple[Path, float, int]]:
    if target_path.is_file():
        st = target_path.stat()
        return [(target_path, st.st_mtime, st.st_size)]
    files: List[Tuple[Path, float, int]] = []
    stack = [str(target_path)]
    while stack:
        curr = stack.pop()
        try:
            with os.scandir(curr) as it:
                for entry in it:
                    name = entry.name
                    if name.startswith("."):
                        continue
                    if entry.is_dir(follow_symlinks=False):
                        if name.endswith(".app") or name in IGNORED_DIRS:
                            continue
                        stack.append(entry.path)
                    elif entry.is_file(follow_symlinks=False):
                        dot = name.rfind(".")
                        if dot != -1 and name[dot:].lower() in ALL_EXTENSIONS:
                            try:
                                st = entry.stat(follow_symlinks=False)
                                files.append((Path(entry.path), st.st_mtime, st.st_size))
                            except Exception:
                                pass
        except (PermissionError, FileNotFoundError):
            continue
    return files

def format_size(bytes_val: int) -> str:
    if bytes_val < 1024:
        return f"{bytes_val} B"
    elif bytes_val < 1024 * 1024:
        return f"{bytes_val / 1024:.1f} KB"
    elif bytes_val < 1024 * 1024 * 1024:
        return f"{bytes_val / (1024 * 1024):.1f} MB"
    return f"{bytes_val / (1024 * 1024 * 1024):.1f} GB"

def get_file_metadata(path: Path, size_bytes: Optional[int] = None, modified_ts: Optional[float] = None) -> Dict[str, Any]:
    if size_bytes is None or modified_ts is None:
        stat = path.stat()
        size_bytes = stat.st_size
        modified_ts = stat.st_mtime
        created_ts = getattr(stat, "st_birthtime", stat.st_ctime)
    else:
        created_ts = modified_ts
    size_human = format_size(size_bytes)
    created_str = datetime.datetime.fromtimestamp(created_ts).strftime("%Y-%m-%d %H:%M:%S")
    modified_str = datetime.datetime.fromtimestamp(modified_ts).strftime("%Y-%m-%d %H:%M:%S")
    ext = path.suffix.lower().lstrip(".")
    return {
        "file_type": ext,
        "size_bytes": size_bytes,
        "size_human": size_human,
        "created_at": created_ts,
        "created_str": created_str,
        "modified_at": modified_ts,
        "modified_str": modified_str,
    }

def extract_ocr_text(path: Path) -> str:
    try:
        import Quartz
        import Vision
        url = Quartz.CFURLCreateWithFileSystemPath(None, str(path.resolve()), Quartz.kCFURLPOSIXPathStyle, False)
        if not url:
            return ""
        request = Vision.VNRecognizeTextRequest.alloc().init()
        request.setRecognitionLevel_(Vision.VNRequestTextRecognitionLevelFast)
        request.setUsesLanguageCorrection_(True)
        handler = Vision.VNImageRequestHandler.alloc().initWithURL_options_(url, None)
        handler.performRequests_error_([request], None)
        lines = []
        for obs in request.results() or []:
            top = obs.topCandidates_(1)
            if top:
                lines.append(top[0].string())
        return " ".join(lines)
    except Exception:
        return ""

def build_photo_record(path: Path, vec: List[float]) -> Dict[str, Any]:
    file_meta = get_file_metadata(path)
    desc = f"Photo: {path.name} | Type: {file_meta['file_type']} | Size: {file_meta['size_human']} | Added/Modified: {file_meta['modified_str']}"
    ocr = extract_ocr_text(path)
    if ocr:
        desc += f"\nOn-screen text: {ocr}"
        file_meta["ocr"] = ocr[:500]
    return {
        "id": hashlib.sha256(f"img:{path}".encode()).hexdigest()[:16],
        "vector": vec,
        "type": "photo",
        "path": str(path.resolve()),
        "filename": path.name,
        "content": desc,
        "modified_at": file_meta["modified_at"],
        "metadata": file_meta,
    }

REPO_MARKERS = (".git", "package.json", "pyproject.toml")

def embed_and_store(emb: Any, store: Any, doc_records: List[Dict[str, Any]], source: str, batch_size: int = 32) -> int:
    if not doc_records:
        return 0
    unembedded = [r for r in doc_records if "vector" not in r]
    if unembedded:
        texts = [r["content"] for r in unembedded]
        titles = [r.get("repo_path") or r.get("filename", "") for r in unembedded]
        vectors = emb.embed_texts(texts, titles=titles, batch_size=batch_size)
        for r, vec in zip(unembedded, vectors):
            r["vector"] = vec
    for r in doc_records:
        r["source"] = source
    return store.add_items(doc_records)

@functools.lru_cache(maxsize=4096)
def _repo_root(directory: Path) -> Optional[Path]:
    if directory == Path.home() or directory == directory.parent:
        return None
    if any((directory / m).exists() for m in REPO_MARKERS):
        return directory
    return _repo_root(directory.parent)

def find_repo_relative_path(path: Path) -> str:
    root = _repo_root(path.parent)
    if root:
        return f"{root.name}/{path.relative_to(root)}"
    return f"{path.parent.name}/{path.name}"

CODE_DEFINERS = (
    "def ", "async def ", "class ", "fn ", "fun ", "func ", "function ",
    "struct ", "enum ", "protocol ", "interface ", "type ", "trait ", "impl ",
    "export default function ", "export function ", "export class "
)

def chunk_code_text(text: str, max_chars: int = 1500) -> List[Tuple[str, int, Optional[str]]]:
    lines = text.splitlines()
    if not lines:
        return []
    hard_split: List[str] = []
    for line in lines:
        while len(line) > max_chars:
            hard_split.append(line[:max_chars])
            line = line[max_chars:]
        hard_split.append(line)
    lines = hard_split

    pieces: List[Tuple[str, int, Optional[str]]] = []
    current_lines: List[str] = []
    start_line = 1
    current_chars = 0
    primary_def: Optional[str] = None

    for idx, line in enumerate(lines, start=1):
        stripped = line.strip()
        is_def = stripped.startswith(CODE_DEFINERS)
        def_label = stripped[:80] if is_def else None
        if is_def and primary_def is None and not line.startswith(" "):
            primary_def = def_label

        line_len = len(line) + 1
        if current_chars + line_len > max_chars and current_lines:
            pieces.append(("\n".join(current_lines), start_line, primary_def))
            overlap = current_lines[-2:] if not is_def and len(current_lines) > 2 else []
            current_lines = list(overlap)
            start_line = idx - len(overlap)
            current_chars = sum(len(l) + 1 for l in current_lines)
            primary_def = def_label

        current_lines.append(line)
        current_chars += line_len

    if current_lines:
        pieces.append(("\n".join(current_lines), start_line, primary_def))
    return pieces

SF_DATALESS = 0x40000000

def is_dataless(path: Path) -> bool:
    """Cloud-only file (e.g. Google Drive streaming); reading its content would download it."""
    try:
        return bool(os.stat(path).st_flags & SF_DATALESS)
    except (OSError, AttributeError):
        return False

def metadata_record(path: Path, record_type: str, label: str, id_prefix: str, **extra_meta: Any) -> Dict[str, Any]:
    """Index a file by name, folder and file details when its content can't or shouldn't be read."""
    file_meta = get_file_metadata(path)
    return {
        "id": hashlib.sha256(f"{id_prefix}:{path}".encode()).hexdigest()[:16],
        "type": record_type,
        "path": str(path),
        "filename": path.name,
        "content": (
            f"{label}: {path.stem}\nFolder: {path.parent}\nType: {file_meta['file_type']}\n"
            f"Size: {file_meta['size_human']}\nModified: {file_meta['modified_str']}"
        ),
        "modified_at": file_meta["modified_at"],
        "metadata": {**file_meta, **extra_meta},
    }

def cloud_placeholder_record(path: Path) -> Dict[str, Any]:
    return metadata_record(path, "document", "File", "cloud", cloud_only=True)

def email_record(record_id: str, source: str, path: str, subject: str, sender: str, to: str, date: str, body: str, modified_at: float) -> Dict[str, Any]:
    return {
        "id": record_id,
        "type": "email",
        "path": path,
        "filename": subject,
        "source": source,
        "content": f"Email: {subject}\nFrom: {sender}\nTo: {to}\nDate: {date}\n\n{body}",
        "modified_at": modified_at,
        "metadata": {"from": sender, "to": to, "date": date, "subject": subject},
    }

def _email_body_text(msg: email.message.EmailMessage) -> str:
    part = msg.get_body(preferencelist=("plain", "html"))
    if part is None:
        return ""
    text = part.get_content()
    if part.get_content_type() == "text/html":
        text = re.sub(r"<(script|style)[^>]*>.*?</\1>|<[^>]+>", " ", text, flags=re.S | re.I)
        text = html.unescape(text)
    return " ".join(text.split())

class ProgressReporter:
    """Throttles progress callbacks to whole-percent steps plus the last file.

    The UI percent ticks 1, 2, 3... as files stream by, but the callback
    (a lock + status string in the job manager) runs at most ~101 times
    per scan instead of once per file.
    """

    def __init__(self, callback: Optional[Callable[[str, int, int], None]], total: int) -> None:
        self.callback = callback
        self.total = total
        self.last_percent = -1

    def report(self, filename: str, processed: int) -> None:
        if self.callback is None or self.total <= 0:
            return
        percent = int(processed * 100 / self.total)
        if percent != self.last_percent or processed >= self.total:
            self.last_percent = percent
            self.callback(filename, processed, self.total)


class Scanner:
    def __init__(self, emb: embedder.Embedder, store: storage.Storage):
        self.embedder = emb
        self.storage = store

    def parse_pdf(self, path: Path) -> List[Dict[str, Any]]:
        chunks = []
        try:
            file_meta = get_file_metadata(path)
            reader = pypdf.PdfReader(str(path))
            pdf = None
            for page_num, page in enumerate(reader.pages):
                text = (page.extract_text() or "").strip()
                if len(text) < 100:
                    try:
                        if pdf is None:
                            pdf = pdfium.PdfDocument(str(path))
                        vec = self.embedder.embed_image(pdf[page_num].render(scale=1.0).to_pil())
                    except Exception:
                        vec = None
                    if vec:
                        chunks.append({
                            "id": hashlib.sha256(f"{path}:pimg:{page_num}".encode()).hexdigest()[:16],
                            "vector": vec,
                            "type": "document",
                            "path": str(path.resolve()),
                            "filename": path.name,
                            "content": f"Document Page Image: {path.name} (Page {page_num + 1})",
                            "modified_at": file_meta["modified_at"],
                            "metadata": {**file_meta, "page": page_num + 1, "is_scanned_image": True},
                        })
                        continue
                if not text:
                    continue
                page_chunks = self.chunk_text(text, chunk_size=800, overlap=100)
                for chunk_idx, chunk in enumerate(page_chunks):
                    chunk_id = hashlib.sha256(f"{path}:{page_num}:{chunk_idx}".encode()).hexdigest()[:16]
                    chunk_meta = dict(file_meta)
                    chunk_meta.update({"page": page_num + 1, "chunk_index": chunk_idx})
                    chunks.append({
                        "id": chunk_id,
                        "type": "document",
                        "path": str(path.resolve()),
                        "filename": path.name,
                        "content": chunk,
                        "modified_at": file_meta["modified_at"],
                        "metadata": chunk_meta,
                    })
        except Exception:
            return []
        return chunks

    def parse_docx(self, path: Path) -> List[Dict[str, Any]]:
        chunks = []
        try:
            file_meta = get_file_metadata(path)
            doc = docx.Document(str(path))
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            if not full_text:
                return []
            text_chunks = self.chunk_text(full_text, chunk_size=800, overlap=100)
            for chunk_idx, chunk in enumerate(text_chunks):
                chunk_id = hashlib.sha256(f"{path}:docx:{chunk_idx}".encode()).hexdigest()[:16]
                chunk_meta = dict(file_meta)
                chunk_meta.update({"chunk_index": chunk_idx})
                chunks.append({
                    "id": chunk_id,
                    "type": "document",
                    "path": str(path.resolve()),
                    "filename": path.name,
                    "content": chunk,
                    "modified_at": file_meta["modified_at"],
                    "metadata": chunk_meta,
                })
        except Exception:
            return []
        return chunks

    def parse_text_file(self, path: Path) -> List[Dict[str, Any]]:
        chunks = []
        try:
            file_meta = get_file_metadata(path)
            if file_meta["size_bytes"] > 2 * 1024 * 1024:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read(500 * 1024).strip()
            else:
                content = path.read_text(encoding="utf-8", errors="ignore").strip()
            if not content:
                return []
            text_chunks = self.chunk_text(content, chunk_size=800, overlap=100)
            for chunk_idx, chunk in enumerate(text_chunks):
                chunk_id = hashlib.sha256(f"{path}:txt:{chunk_idx}".encode()).hexdigest()[:16]
                chunk_meta = dict(file_meta)
                chunk_meta.update({"chunk_index": chunk_idx})
                chunks.append({
                    "id": chunk_id,
                    "type": "document",
                    "path": str(path.resolve()),
                    "filename": path.name,
                    "content": chunk,
                    "modified_at": file_meta["modified_at"],
                    "metadata": chunk_meta,
                })
        except Exception:
            return []
        return chunks

    def parse_code_file(self, path: Path) -> List[Dict[str, Any]]:
        chunks = []
        try:
            file_meta = get_file_metadata(path)
            if file_meta["size_bytes"] > 1024 * 1024:
                return []
            content = path.read_text(encoding="utf-8", errors="ignore")
            top = content[:1000]
            if "PRIVATE KEY-----" in top:
                return []
            if any(mark in top for mark in ("DO NOT EDIT", "@generated", "Code generated by")):
                return []

            repo_rel = find_repo_relative_path(path)
            pieces = chunk_code_text(content, max_chars=1500)
            for idx, (piece_text, first_line, def_name) in enumerate(pieces):
                chunk_id = hashlib.sha256(f"{path}:code:{idx}".encode()).hexdigest()[:16]
                cmeta = dict(file_meta)
                cmeta.update({
                    "repo_path": repo_rel,
                    "first_line": first_line,
                    "definition": def_name or "",
                    "is_code": True,
                })
                chunks.append({
                    "id": chunk_id,
                    "type": "document",
                    "path": str(path.resolve()),
                    "filename": path.name,
                    "content": piece_text,
                    "modified_at": file_meta["modified_at"],
                    "metadata": cmeta,
                    "repo_path": repo_rel,
                })
        except Exception:
            return []
        return chunks

    def parse_image(self, path: Path) -> Optional[Dict[str, Any]]:
        vec = self.embedder.embed_image(path)
        if vec is None:
            return None
        return build_photo_record(path, vec)

    def parse_email(self, path: Path) -> List[Dict[str, Any]]:
        try:
            raw = path.read_bytes()
            if path.suffix.lower() == ".emlx":
                length_line, _, rest = raw.partition(b"\n")
                raw = rest[: int(length_line.strip())]
            msg = email.message_from_bytes(raw, policy=email.policy.default)
            subject = str(msg.get("Subject") or "(no subject)")
            return [email_record(
                record_id=hashlib.sha256(f"mail:{path}".encode()).hexdigest()[:16],
                source="apple_mail",
                path=str(path.resolve()),
                subject=subject,
                sender=str(msg.get("From") or ""),
                to=str(msg.get("To") or ""),
                date=str(msg.get("Date") or ""),
                body=_email_body_text(msg)[:2000],
                modified_at=path.stat().st_mtime,
            )]
        except Exception:
            return []

    def parse_vcf(self, path: Path) -> List[Dict[str, Any]]:
        records = []
        try:
            content = path.read_text(encoding="utf-8", errors="ignore")
            file_meta = get_file_metadata(path)
            lines = content.splitlines()
            current_contact: Dict[str, str] = {}
            for line in lines:
                line = line.strip()
                if line.startswith("BEGIN:VCARD"):
                    current_contact = {}
                elif line.startswith("FN:"):
                    current_contact["name"] = line[3:].strip()
                elif line.startswith("TEL"):
                    current_contact["phone"] = line.split(":", 1)[-1].strip()
                elif line.startswith("EMAIL"):
                    current_contact["email"] = line.split(":", 1)[-1].strip()
                elif line.startswith("ORG:"):
                    current_contact["org"] = line[4:].strip()
                elif line.startswith("TITLE:"):
                    current_contact["title"] = line[6:].strip()
                elif line.startswith("NOTE:"):
                    current_contact["note"] = line[5:].strip()
                elif line.startswith("END:VCARD"):
                    if current_contact.get("name"):
                        contact_str = f"Contact: {current_contact['name']}"
                        if current_contact.get("org"):
                            contact_str += f" | Org: {current_contact['org']}"
                        if current_contact.get("title"):
                            contact_str += f" | Title: {current_contact['title']}"
                        if current_contact.get("phone"):
                            contact_str += f" | Phone: {current_contact['phone']}"
                        if current_contact.get("email"):
                            contact_str += f" | Email: {current_contact['email']}"
                        if current_contact.get("note"):
                            contact_str += f" | Note: {current_contact['note']}"
                        cid = hashlib.sha256(f"vcf:{path}:{current_contact['name']}".encode()).hexdigest()[:16]
                        cmeta = dict(file_meta)
                        cmeta.update(current_contact)
                        records.append({
                            "id": cid,
                            "type": "contact",
                            "path": str(path.resolve()),
                            "filename": path.name,
                            "content": contact_str,
                            "modified_at": file_meta["modified_at"],
                            "metadata": cmeta,
                        })
        except Exception:
            return []
        return records

    def parse_ics(self, path: Path) -> List[Dict[str, Any]]:
        records = []
        try:
            content = path.read_text(encoding="utf-8", errors="ignore")
            file_meta = get_file_metadata(path)
            lines = content.splitlines()
            current_event: Dict[str, str] = {}
            for line in lines:
                line = line.strip()
                if line.startswith("BEGIN:VEVENT"):
                    current_event = {}
                elif line.startswith("SUMMARY:"):
                    current_event["summary"] = line[8:].strip()
                elif line.startswith("DTSTART"):
                    current_event["start"] = line.split(":", 1)[-1].strip()
                elif line.startswith("DTEND"):
                    current_event["end"] = line.split(":", 1)[-1].strip()
                elif line.startswith("LOCATION:"):
                    current_event["location"] = line[9:].strip()
                elif line.startswith("DESCRIPTION:"):
                    current_event["description"] = line[12:].strip()
                elif line.startswith("END:VEVENT"):
                    if current_event.get("summary"):
                        event_str = f"Calendar Event: {current_event['summary']}"
                        if current_event.get("start"):
                            event_str += f" | Start: {current_event['start']}"
                        if current_event.get("end"):
                            event_str += f" | End: {current_event['end']}"
                        if current_event.get("location"):
                            event_str += f" | Location: {current_event['location']}"
                        if current_event.get("description"):
                            event_str += f" | Details: {current_event['description']}"
                        eid = hashlib.sha256(f"ics:{path}:{current_event['summary']}".encode()).hexdigest()[:16]
                        emeta = dict(file_meta)
                        emeta.update(current_event)
                        records.append({
                            "id": eid,
                            "type": "calendar",
                            "path": str(path.resolve()),
                            "filename": path.name,
                            "content": event_str,
                            "modified_at": file_meta["modified_at"],
                            "metadata": emeta,
                        })
        except Exception:
            return []
        return records

    def parse_bookmark(self, path: Path) -> List[Dict[str, Any]]:
        try:
            content = path.read_text(encoding="utf-8", errors="ignore")
            file_meta = get_file_metadata(path)
            url = ""
            for line in content.splitlines():
                if "URL=" in line:
                    url = line.split("URL=", 1)[-1].strip()
                    break
                elif "<string>http" in line:
                    url = line.replace("<string>", "").replace("</string>", "").strip()
                    break
            b_str = f"Web Bookmark: {path.stem} | URL: {url}"
            bid = hashlib.sha256(f"bm:{path}".encode()).hexdigest()[:16]
            bmeta = dict(file_meta)
            bmeta["url"] = url
            return [{
                "id": bid,
                "type": "bookmark",
                "path": str(path.resolve()),
                "filename": path.name,
                "content": b_str,
                "modified_at": file_meta["modified_at"],
                "metadata": bmeta,
            }]
        except Exception:
            return []

    def chunk_text(self, text: str, chunk_size: int = 800, overlap: int = 100) -> List[str]:
        if len(text) <= chunk_size:
            return [text]
        chunks = []
        start = 0
        while start < len(text):
            end = start + chunk_size
            chunks.append(text[start:end])
            start += chunk_size - overlap
        return chunks

    def scan_path(
        self,
        target_path: Path,
        source: str = "local",
        on_progress: Optional[Callable[[str, int, int], None]] = None,
        cancel_check: Optional[Callable[[], bool]] = None,
    ) -> int:
        if not target_path.exists():
            return 0

        files_to_process = collect_supported_files(target_path)
        total_files = len(files_to_process)
        if total_files == 0:
            return 0

        indexed_mtimes = self.storage.get_indexed_file_mtimes(under=str(target_path))
        indexed_items = 0

        doc_buffer: List[Dict[str, Any]] = []
        image_buffer: List[Path] = []
        BATCH_SIZE = 64
        IMG_BATCH_SIZE = 16
        MAX_QUEUED_BATCHES = 4

        embed_pool = ThreadPoolExecutor(max_workers=1)
        pending: List[Future] = []

        def submit_embedding(fn, *args):
            # Parsing outruns embedding by a wide margin; without a bound the
            # progress bar hits 100% and every parsed batch piles up in memory.
            while True:
                unfinished = [f for f in pending if not f.done()]
                if len(unfinished) < MAX_QUEUED_BATCHES:
                    break
                wait(unfinished, return_when=FIRST_COMPLETED)
            pending.append(embed_pool.submit(fn, *args))

        def flush_docs():
            nonlocal doc_buffer
            if not doc_buffer:
                return
            batch, doc_buffer = doc_buffer, []
            submit_embedding(embed_and_store, self.embedder, self.storage, batch, source, BATCH_SIZE)

        def flush_images():
            nonlocal image_buffer
            if not image_buffer:
                return
            batch, image_buffer = image_buffer, []
            submit_embedding(self._embed_image_batch, batch, source, IMG_BATCH_SIZE)

        workers = min(8, os.cpu_count() or 4)
        CHUNK_SIZE = 64
        reporter = ProgressReporter(on_progress, total_files)
        processed_files = 0

        # One pool per scan, not per chunk: creating and tearing down a pool
        # for every 64 files churns threads for no reason.
        pool = ThreadPoolExecutor(max_workers=workers)
        cancelled = False
        try:
            try:
                for chunk_start in range(0, total_files, CHUNK_SIZE):
                    if cancel_check and cancel_check():
                        cancelled = True
                        break

                    chunk_files = files_to_process[chunk_start : chunk_start + CHUNK_SIZE]
                    files_to_parse = []
                    for file_path, st_mtime, st_size in chunk_files:
                        try:
                            resolved_str = str(file_path.resolve())
                            if resolved_str in indexed_mtimes and abs(indexed_mtimes[resolved_str] - st_mtime) < 0.01:
                                continue
                        except Exception:
                            pass
                        files_to_parse.append(file_path)

                    # Already-indexed files count as processed without parsing.
                    processed_files += len(chunk_files) - len(files_to_parse)
                    if not files_to_parse:
                        reporter.report(chunk_files[-1][0].name, processed_files)
                        continue

                    # map streams results in order so progress and embedding
                    # advance file-by-file instead of waiting for the whole chunk.
                    for kind, fp, chunks in pool.map(self._parse_file, files_to_parse):
                        if kind == "image":
                            image_buffer.append(fp)
                            if len(image_buffer) >= IMG_BATCH_SIZE:
                                flush_images()
                        elif kind == "doc" and chunks:
                            doc_buffer.extend(chunks)
                            if len(doc_buffer) >= BATCH_SIZE:
                                flush_docs()
                        processed_files += 1
                        reporter.report(fp.name, processed_files)
            finally:
                pool.shutdown(wait=True)

            if not cancelled:
                flush_docs()
                flush_images()
            total_batches = len(pending)
            for i, future in enumerate(pending):
                if not cancelled and cancel_check and cancel_check():
                    cancelled = True
                if cancelled:
                    # Queued batches must not write once the scan is stopped
                    # (a Reset would otherwise be repopulated behind its back).
                    embed_pool.shutdown(wait=True, cancel_futures=True)
                if future.cancelled():
                    continue
                if total_batches > 1:
                    reporter.report(f"Embedding batch {i + 1}/{total_batches}...", processed_files)
                indexed_items += future.result()
        finally:
            embed_pool.shutdown(wait=True, cancel_futures=True)

        return indexed_items

    def _embed_image_batch(self, image_buffer: List[Path], source: str, batch_size: int) -> int:
        vectors = self.embedder.embed_images(image_buffer, batch_size=batch_size)
        records = []
        for img_path, vec in zip(image_buffer, vectors):
            if vec is None:
                continue
            records.append({**build_photo_record(img_path, vec), "source": source})
        if records:
            return self.storage.add_items(records)
        return 0

    def _parse_file(self, file_path: Path) -> Tuple[str, Path, List[Dict[str, Any]]]:
        ext = file_path.suffix.lower()
        if is_dataless(file_path):
            return ("doc", file_path, [cloud_placeholder_record(file_path)])
        if ext in IMAGE_EXTENSIONS:
            return ("image", file_path, [])
        elif ext == ".pdf":
            return ("doc", file_path, self.parse_pdf(file_path))
        elif ext == ".docx":
            return ("doc", file_path, self.parse_docx(file_path))
        elif ext in CONTACT_EXTENSIONS:
            return ("doc", file_path, self.parse_vcf(file_path))
        elif ext in CALENDAR_EXTENSIONS:
            return ("doc", file_path, self.parse_ics(file_path))
        elif ext in BOOKMARK_EXTENSIONS:
            return ("doc", file_path, self.parse_bookmark(file_path))
        elif ext in EMAIL_EXTENSIONS:
            return ("doc", file_path, self.parse_email(file_path))
        elif ext in CODE_EXTENSIONS:
            return ("doc", file_path, self.parse_code_file(file_path))
        elif ext in TEXT_EXTENSIONS:
            return ("doc", file_path, self.parse_text_file(file_path))
        elif ext in VIDEO_EXTENSIONS:
            return ("doc", file_path, [metadata_record(file_path, "video", "Video", "media")])
        elif ext in AUDIO_EXTENSIONS:
            return ("doc", file_path, [metadata_record(file_path, "audio", "Audio", "media")])
        return ("unknown", file_path, [])

